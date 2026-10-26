"""Produce and verify synthetic Office snapshots through the real application API.

Copyright 2026 Dhruba Poudel
SPDX-License-Identifier: Apache-2.0
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import platform
import tempfile
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from figurerelay.app import create_app
from figurerelay.exports import AUTHOR


def verify_demo(output: Path) -> dict:
    """Exercise two revisions without using or modifying normal application data."""
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        raise ValueError("Choose a new or empty output directory; existing files are never replaced.")
    checks: list[str] = []

    def require(condition: bool, description: str) -> None:
        if not condition:
            raise RuntimeError(description)
        checks.append(description)

    def request(client, method, path, status=200, **kwargs):
        response = client.request(method, path, **kwargs)
        if response.status_code != status:
            raise RuntimeError(f"{method} {path}: expected {status}, got {response.status_code}: {response.text}")
        return response

    def inspect_bundle(data: bytes, generation: dict, expected: str, label: str) -> dict:
        with ZipFile(BytesIO(data)) as archive:
            require(set(archive.namelist()) == {"report.docx", "slides.pptx", "manifest.json"},
                    f"{label}: ZIP contains the three expected files")
            files = {name: archive.read(name) for name in archive.namelist()}
        report = Document(BytesIO(files["report.docx"]))
        slides = Presentation(BytesIO(files["slides.pptx"]))
        report_text = "\n".join(paragraph.text for paragraph in report.paragraphs)
        slide_text = "\n".join(shape.text for slide in slides.slides for shape in slide.shapes
                               if shape.has_text_frame)
        require(report_text.count(expected) == 2 and slide_text.count(expected) == 2,
                f"{label}: the expected revenue appears in all four Office occurrences")
        for kind, properties in (("report", report.core_properties), ("slides", slides.core_properties)):
            require(properties.author == AUTHOR and properties.last_modified_by == AUTHOR,
                    f"{label}: {kind} authorship is Dhruba Poudel")
            require(properties.identifier == generation["id"],
                    f"{label}: {kind} identifies the reviewed generation")
        manifest = json.loads(files["manifest.json"])
        require(manifest["generation_id"] == generation["id"]
                and manifest["revision"] == generation["revision"],
                f"{label}: manifest identifies the reviewed revision")
        require({item["name"] for item in manifest["artifacts"]} == {"report.docx", "slides.pptx"},
                f"{label}: manifest lists both Office artifacts")
        for artifact in manifest["artifacts"]:
            content = files[artifact["name"]]
            require(artifact["sha256"] == hashlib.sha256(content).hexdigest()
                    and artifact["size"] == len(content),
                    f"{label}: {artifact['name']} checksum and size match the exported bytes")
        return {"generation_id": generation["id"], "revision": generation["revision"],
                "expected_revenue": expected, "zip_sha256": hashlib.sha256(data).hexdigest(),
                "files": files}

    with tempfile.TemporaryDirectory(prefix="figurerelay-demo-") as scratch:
        with TestClient(create_app(Path(scratch)), base_url="http://127.0.0.1:8080") as client:
            project = request(client, "POST", "/api/demo", 201).json()
            base = f"/api/projects/{project['id']}"
            first = request(client, "POST", base + "/preview", 201,
                            json={"expected_revision": project["revision"]}).json()
            require(not first["errors"], "v1: synthetic source produces a valid candidate")
            first_url = base + f"/generations/{first['id']}"
            request(client, "GET", first_url + "/export", 409)
            checks.append("v1: export is blocked before exact-revision approval")
            request(client, "POST", first_url + "/approve", json={
                "expected_revision": project["revision"], "actor": AUTHOR, "acknowledge_warnings": True})
            first_zip = request(client, "GET", first_url + "/export").content
            first_result = inspect_bundle(first_zip, first, "1,250,000 AUD", "v1")

            current = request(client, "POST", base + "/demo-revision").json()
            second = request(client, "POST", base + "/preview", 201,
                             json={"expected_revision": current["revision"]}).json()
            require(not second["errors"], "v2: changed synthetic source produces a valid candidate")
            changes = [change for change in second["changes"] if change["key"] == "revenue"]
            require(len(changes) == 4 and all(change["before"] == "1,250,000 AUD"
                    and change["after"] == "1,420,000 AUD" for change in changes),
                    "v2: review identifies the same revenue change in all four linked occurrences")
            second_url = base + f"/generations/{second['id']}"
            request(client, "POST", second_url + "/approve", json={
                "expected_revision": current["revision"], "actor": AUTHOR, "acknowledge_warnings": True})
            second_zip = request(client, "GET", second_url + "/export").content
            second_result = inspect_bundle(second_zip, second, "1,420,000 AUD", "v2")
            require(request(client, "GET", first_url + "/export").content == first_zip,
                    "v1: approved ZIP remains byte-identical after v2 is approved")

        with TestClient(create_app(Path(scratch)), base_url="http://127.0.0.1:8080") as restarted:
            require(request(restarted, "GET", first_url + "/export").content == first_zip,
                    "v1: approved ZIP remains byte-identical after application restart")

    output.mkdir(parents=True, exist_ok=True)
    bundles = []
    for label, data, result in (("v1", first_zip, first_result), ("v2", second_zip, second_result)):
        destination = output / label
        destination.mkdir()
        for name, content in result.pop("files").items():
            (destination / name).write_bytes(content)
        (output / f"figurerelay-demo-{label}.zip").write_bytes(data)
        bundles.append({"version": label, **result})
    summary = {"author": AUTHOR, "last_modified_by": AUTHOR,
               "verified_at_utc": datetime.now(timezone.utc).isoformat(),
               "python": platform.python_version(), "status": "passed", "data": "built-in synthetic demo",
               "checks": checks, "bundles": bundles,
               "limits": ["Office text, metadata and checksums were inspected; visual layout was not certified.",
                          "This local single-user check does not validate organisational identities or immutable audit storage."]}
    (output / "verification.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="New or empty directory for synthetic exports and verification.json")
    options = parser.parse_args()
    try:
        summary = verify_demo(options.output)
    except (ValueError, RuntimeError) as error:
        parser.exit(1, f"Verification failed: {error}\n")
    print(f"Passed {len(summary['checks'])} workflow and artifact checks. Output: {options.output.resolve()}")


if __name__ == "__main__":
    main()
