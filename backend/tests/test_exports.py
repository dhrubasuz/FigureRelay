"""Validate readable exports, provenance, and requested authorship.

Copyright 2026 Dhruba Poudel
SPDX-License-Identifier: Apache-2.0
"""
import hashlib
from io import BytesIO
import json

from docx import Document
from pptx import Presentation
import pytest

from figurerelay.exports import AUTHOR, build_exports


def rendered():
    return {
        "title": "Quarterly report — Café",
        "sections": [{"title": "Revenue", "body": "Revenue was USD 12,400,000.\nZero is 0."}],
        "slides": [{"title": "Quarterly review", "body": "Revenue was USD 12,400,000."}],
    }


def test_office_snapshots_and_authorship():
    files = build_exports(rendered(), {"generation_id": "gen-123", "source_hash": "abc"})
    report = Document(BytesIO(files["report.docx"]))
    deck = Presentation(BytesIO(files["slides.pptx"]))
    assert "Revenue was USD 12,400,000." in "\n".join(p.text for p in report.paragraphs)
    assert "Revenue was USD 12,400,000." in "\n".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    for properties in (report.core_properties, deck.core_properties):
        assert properties.author == AUTHOR
        assert properties.last_modified_by == AUTHOR
        assert properties.identifier == "gen-123"
        assert properties.title == rendered()["title"]
    manifest = json.loads(files["manifest.json"])
    assert manifest["export_author"] == AUTHOR
    for item in manifest["artifacts"]:
        assert item["sha256"] == hashlib.sha256(files[item["name"]]).hexdigest()
        assert item["size"] == len(files[item["name"]])


def test_manifest_input_is_not_mutated():
    original = {"generation_id": "gen-123", "sources": [{"hash": "abc"}]}
    build_exports(rendered(), original)
    assert original == {"generation_id": "gen-123", "sources": [{"hash": "abc"}]}


def test_layout_rejects_excessive_slide_text():
    content = rendered()
    content["slides"][0]["body"] = "x" * 1201
    with pytest.raises(ValueError, match="1200"):
        build_exports(content, {})


def test_many_short_lines_do_not_silently_overflow():
    content = rendered()
    content["slides"][0]["body"] = "line\n" * 80
    with pytest.raises(ValueError, match="Split"):
        build_exports(content, {})
