"""Create frozen, attributed Office snapshots from resolved reporting content.

Copyright 2026 Dhruba Poudel
SPDX-License-Identifier: Apache-2.0
"""

from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import textwrap

from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

AUTHOR = "Dhruba Poudel"
INK = (24, 45, 43)
TEAL = (21, 113, 99)


def _text(value: object, limit: int, field: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{field} must be text.")
    if len(value) > limit:
        raise ValueError(f"{field} is limited to {limit} characters.")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise ValueError(f"{field} contains unsupported control characters.")
    return value


def _metadata(properties, title: str, generation: str, timestamp: datetime) -> None:
    properties.author = AUTHOR
    properties.last_modified_by = AUTHOR
    properties.title = title[:255]
    properties.subject = "Approved reporting snapshot"
    properties.keywords = "FigureRelay, linked reporting, snapshot"
    properties.identifier = generation[:255]
    properties.created = timestamp
    properties.modified = timestamp
    properties.comments = "Static snapshot; refresh linked values in FigureRelay."


def _report(rendered: dict, generation: str, timestamp: datetime) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = DocxInches(8.27)
    section.page_height = DocxInches(11.69)
    section.top_margin = section.bottom_margin = DocxInches(0.8)
    section.left_margin = section.right_margin = DocxInches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = DocxPt(11)
    normal.font.color.rgb = DocxColor(*INK)
    normal.paragraph_format.space_after = DocxPt(8)
    for name in ("Title", "Heading 1"):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.color.rgb = DocxColor(*TEAL)
    document.add_heading(rendered["title"], 0)
    document.add_paragraph(f"Prepared by {AUTHOR}", "Subtitle")
    for item in rendered["sections"]:
        document.add_heading(item["title"], 1)
        for line in item["body"].split("\n"):
            document.add_paragraph(line)
    document.add_heading("Source and version", 1)
    document.add_paragraph(f"FigureRelay generation: {generation}")
    document.add_paragraph(
        "This report is a frozen snapshot. The accompanying manifest records its "
        "source revisions and file checksums. Changes require a new review in FigureRelay."
    )
    footer = section.footer.paragraphs[0]
    footer.text = "FigureRelay | Reporting snapshot"
    footer.style = document.styles["Caption"]
    _metadata(document.core_properties, rendered["title"], generation, timestamp)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _textbox(slide, x: float, y: float, width: float, height: float,
             content: str, size: int, color: tuple[int, int, int], bold: bool = False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor(*color)
        paragraph.space_after = Pt(5)
    return box


def _deck(rendered: dict, generation: str, timestamp: datetime) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    items = rendered["slides"] or [{"title": rendered["title"], "body": ""}]
    for index, item in enumerate(items, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(249, 248, 243)
        _textbox(slide, 0.7, 0.35, 11.9, 0.3, "FIGURERELAY / REPORTING SNAPSHOT", 10, TEAL, True)
        title_size = 30
        for proposed_title in (30, 26, 24):
            title_lines = sum(max(1, len(textwrap.wrap(line, width=int(1450 / proposed_title))))
                              for line in item["title"].split("\n"))
            if title_lines * (proposed_title * 1.15 + 5) <= 68:
                title_size = proposed_title
                break
        else:
            raise ValueError("Slide title exceeds the built-in layout. Use a shorter title.")
        _textbox(slide, 0.7, 0.9, 11.9, 1.0, item["title"], title_size, INK, True)
        size = 24
        for proposed in (24, 22, 20, 18, 16):
            width = int(1510 / proposed)
            lines = sum(max(1, len(textwrap.wrap(line, width=width)))
                        for line in item["body"].split("\n"))
            if lines * (proposed * 1.2 + 5) <= 310:
                size = proposed
                break
        else:
            raise ValueError("Slide text exceeds the built-in layout. Split it across slides.")
        _textbox(slide, 0.7, 2.0, 11.9, 4.4, item["body"], size, INK)
        _textbox(slide, 0.7, 6.8, 10.7, 0.3,
                 f"Prepared by {AUTHOR} | Version {generation[:12]}", 10, TEAL)
        _textbox(slide, 12.1, 6.8, 0.5, 0.3, str(index), 10, TEAL)
    _metadata(presentation.core_properties, rendered["title"], generation, timestamp)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def build_exports(rendered: dict, manifest: dict) -> dict[str, bytes]:
    """Return DOCX, PPTX, and provenance bytes; approval/export must not regenerate.

    The manifest hashes the two Office files. It intentionally cannot hash itself.
    The caller can separately hash all three files in its generation record.
    """
    title = _text(rendered.get("title"), 200, "Report title")
    sections = rendered.get("sections", [])
    slides = rendered.get("slides", [])
    if not isinstance(sections, list) or not isinstance(slides, list):
        raise ValueError("Report sections and slides must be lists.")
    if len(sections) > 30 or len(slides) > 20:
        raise ValueError("The built-in layouts allow 30 report sections and 20 slides.")
    for item in sections:
        _text(item.get("title"), 200, "Section title")
        _text(item.get("body"), 10000, "Section body")
    for item in slides:
        _text(item.get("title"), 120, "Slide title")
        _text(item.get("body"), 1200, "Slide body")
    timestamp = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
    generation = str(manifest.get("generation_id", manifest.get("id", "snapshot")))
    content = {"title": title, "sections": sections, "slides": slides}
    files = {
        "report.docx": _report(content, generation, timestamp),
        "slides.pptx": _deck(content, generation, timestamp),
    }
    provenance = deepcopy(manifest)
    provenance["schema_version"] = "1.0"
    provenance["export_author"] = AUTHOR
    provenance["snapshot_notice"] = "Frozen Office output. Refresh and approve changes in FigureRelay."
    provenance["artifacts"] = [
        {"name": name, "sha256": hashlib.sha256(data).hexdigest(), "size": len(data)}
        for name, data in files.items()
    ]
    files["manifest.json"] = (json.dumps(provenance, ensure_ascii=False, indent=2,
                                         allow_nan=False) + "\n").encode("utf-8")
    return files
