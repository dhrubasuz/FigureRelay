"""Read a narrow source format without evaluating formulas or rewriting files.

Copyright 2026 Dhruba Poudel
SPDX-License-Identifier: Apache-2.0
"""

import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import io
import re
import zipfile

from openpyxl import load_workbook


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ZIP_MEMBERS = 2000
MAX_METRICS = 5000
HEADERS = ("key", "label", "value", "kind", "unit", "period")
KEY_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9_.:-]{0,119}$")


class SourceError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSource:
    metrics: list[dict]
    warnings: list[str]
    sha256: str


def _headers(row) -> dict[str, int]:
    names = [str(value).strip().lower() if value is not None else "" for value in row]
    if len(names) != len(set(names)):
        raise SourceError("Source has duplicate column headers.")
    if set(names) != set(HEADERS):
        raise SourceError("Source must contain exactly these columns: " + ", ".join(HEADERS) + ".")
    return {name: index for index, name in enumerate(names)}


def _text(value, field: str, row: int, maximum: int = 240) -> str:
    result = "" if value is None else str(value).strip()
    if len(result) > maximum:
        raise SourceError(f"Row {row}: {field} exceeds {maximum} characters.")
    return result


def _metric(values: dict, row: int) -> dict:
    key = _text(values["key"], "key", row, 120)
    if not KEY_PATTERN.fullmatch(key):
        raise SourceError(f"Row {row}: key must start with a letter and contain only letters, numbers, _, ., :, or -.")
    label = _text(values["label"], "label", row)
    if not label:
        raise SourceError(f"Row {row}: label is required.")
    kind = _text(values["kind"], "kind", row, 20).lower()
    raw = values["value"]
    if kind == "number":
        if raw is None or isinstance(raw, bool) or not str(raw).strip():
            raise SourceError(f"Row {row}: number value is required.")
        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", str(raw).strip()):
            raise SourceError(f"Row {row}: number must be a finite decimal without currency symbols or separators.")
        try:
            numeric = Decimal(str(raw).strip())
        except InvalidOperation as exc:
            raise SourceError(f"Row {row}: value must be a decimal number without currency symbols or separators.") from exc
        if not numeric.is_finite() or abs(numeric.adjusted()) > 30 or len(numeric.as_tuple().digits) > 40:
            raise SourceError(f"Row {row}: number must be finite with at most 40 digits and an exponent between -30 and 30.")
        value = format(numeric, "f")
    elif kind == "date":
        try:
            if isinstance(raw, datetime):
                value = raw.date().isoformat()
            elif isinstance(raw, date):
                value = raw.isoformat()
            else:
                text = str(raw).strip()
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
                    raise ValueError("Date is not YYYY-MM-DD")
                value = date.fromisoformat(text).isoformat()
        except (ValueError, TypeError) as exc:
            raise SourceError(f"Row {row}: date must use YYYY-MM-DD.") from exc
    elif kind == "text":
        value = "" if raw is None else str(raw)
        if len(value) > 2000:
            raise SourceError(f"Row {row}: text value exceeds 2000 characters.")
    else:
        raise SourceError(f"Row {row}: kind must be number, text, or date.")
    return {"key": key, "label": label, "value": value, "kind": kind,
            "unit": _text(values["unit"], "unit", row, 40),
            "period": _text(values["period"], "period", row, 120)}


def _append(metrics: list, seen: set, values: dict, row: int) -> None:
    if all(value is None or str(value).strip() == "" for value in values.values()):
        return
    metric = _metric(values, row)
    if metric["key"] in seen:
        raise SourceError(f"Row {row}: duplicate key '{metric['key']}'. Every metric needs a unique key.")
    if len(metrics) >= MAX_METRICS:
        raise SourceError(f"Source exceeds {MAX_METRICS} metrics.")
    seen.add(metric["key"])
    metrics.append(metric)


def _csv(data: bytes) -> tuple[list[dict], list[str]]:
    try:
        content = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SourceError("CSV must use UTF-8 encoding.") from exc
    if "\x00" in content:
        raise SourceError("CSV contains invalid null characters.")
    try:
        reader = csv.reader(io.StringIO(content, newline=""), strict=True)
        first = next(reader, None)
        if first is None:
            raise SourceError("Source is empty.")
        columns = _headers(first)
        metrics, seen = [], set()
        for row_number, row in enumerate(reader, start=2):
            if not row or all(not cell.strip() for cell in row):
                continue
            if len(row) != len(columns):
                raise SourceError(f"Row {row_number}: column count does not match the header.")
            _append(metrics, seen, {name: row[index] for name, index in columns.items()}, row_number)
    except csv.Error as exc:
        raise SourceError("CSV is malformed: " + str(exc)) from exc
    return metrics, []


def _check_xlsx(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            names = [item.filename for item in infos]
            if len(infos) > MAX_ZIP_MEMBERS or sum(item.file_size for item in infos) > MAX_EXPANDED_BYTES:
                raise SourceError("XLSX expanded content exceeds the safe import limit (25 MB / 2000 parts).")
            if len(names) != len(set(names)):
                raise SourceError("XLSX contains duplicate ZIP entries.")
            if any(item.flag_bits & 1 for item in infos):
                raise SourceError("Encrypted workbooks are not supported.")
            if any(name.startswith("/") or ".." in name.split("/") or "\\" in name for name in names):
                raise SourceError("XLSX contains invalid ZIP paths.")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise SourceError("Macro-enabled workbooks are not supported.")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise SourceError("File is not an XLSX workbook.")
    except (zipfile.BadZipFile, NotImplementedError) as exc:
        raise SourceError("File is not a readable XLSX workbook.") from exc


def _xlsx(data: bytes) -> tuple[list[dict], list[str]]:
    _check_xlsx(data)
    formula_book = cached_book = None
    try:
        formula_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        if "Metrics" not in formula_book.sheetnames:
            raise SourceError("XLSX must have a worksheet named Metrics.")
        formula_sheet, cached_sheet = formula_book["Metrics"], cached_book["Metrics"]
        if (formula_sheet.max_row or 0) > MAX_METRICS + 1 or (formula_sheet.max_column or 0) > len(HEADERS):
            raise SourceError("Metrics worksheet exceeds 5000 rows or contains columns outside the six-column schema.")
        rows = formula_sheet.iter_rows()
        cached_rows = cached_sheet.iter_rows()
        first, cached_first = next(rows, None), next(cached_rows, None)
        if first is None:
            raise SourceError("Metrics worksheet is empty.")
        if any(cell.data_type == "f" for cell in first):
            raise SourceError("Column headers must be literal text.")
        columns = _headers([cell.value for cell in first])
        metrics, warnings, seen = [], [], set()
        for row_number, (row, cached_row) in enumerate(zip(rows, cached_rows), start=2):
            values, formula_value = {}, False
            for name, index in columns.items():
                cell, cached = row[index], cached_row[index]
                if cell.data_type == "e" or cached.data_type == "e":
                    raise SourceError(f"Row {row_number}: {name} contains an Excel error.")
                if cell.data_type == "f":
                    if name != "value":
                        raise SourceError(f"Row {row_number}: formulas are supported only in the value column.")
                    if cached.value is None:
                        raise SourceError(f"Row {row_number}: formula has no stored result. Recalculate and save it in Excel before importing.")
                    values[name], formula_value = cached.value, True
                else:
                    values[name] = cell.value
            before_count = len(metrics)
            _append(metrics, seen, values, row_number)
            if formula_value and len(metrics) > before_count:
                warnings.append(f"{metrics[-1]['key']}: imported a stored Excel formula result; its freshness cannot be verified. Recalculate and save the workbook before approving.")
        return metrics, warnings
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError("Unable to read XLSX. Check that the workbook is valid and unencrypted.") from exc
    finally:
        if formula_book is not None:
            formula_book.close()
        if cached_book is not None:
            cached_book.close()


def parse_source(filename: str, data: bytes) -> ParsedSource:
    if not data:
        raise SourceError("Source is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise SourceError("Source exceeds the 5 MB upload limit.")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        if data.startswith(b"PK"):
            raise SourceError("CSV file contains ZIP data.")
        metrics, warnings = _csv(data)
    elif suffix == "xlsx":
        metrics, warnings = _xlsx(data)
    else:
        raise SourceError("Choose a .csv or .xlsx source file.")
    if not metrics:
        raise SourceError("Source contains no metrics.")
    return ParsedSource(metrics, warnings, hashlib.sha256(data).hexdigest())
